"""Описание тендера в Telegram/Excel и лимит 4096 символов Telegram."""
from __future__ import annotations

from openpyxl import load_workbook

from src.export.excel import export_tenders_to_excel
from src.models.tender import Tender, TenderAnalysis
from src.notifications.telegram import TELEGRAM_MAX_MESSAGE, TelegramNotifier
from src.storage.database import TenderDatabase


def _analysis(summary="Ок", risks=None):
    return TenderAnalysis(
        relevance_score=80, summary=summary, risks=risks or [], recommendation="participate",
    )


def test_message_shows_real_description_and_platform():
    t = Tender(platform="b2b_center", external_id="1", title="Станок", url="https://x/1",
               description="Поставка токарного станка с ЧПУ")
    msg = TelegramNotifier.format_message(t, _analysis())
    assert "Описание:" in msg and "токарного станка" in msg
    assert "B2B-Center" in msg


def test_message_has_no_description_line_when_description_empty():
    t = Tender(platform="eis", external_id="1", title="Станок", url="https://x/1")
    assert "Описание:" not in TelegramNotifier.format_message(t, _analysis())


def test_message_never_exceeds_telegram_limit_and_keeps_link_and_valid_html():
    t = Tender(platform="eis", external_id="1", title="Станок", url="https://x/1",
               description="д" * 5000)
    msg = TelegramNotifier.format_message(
        t, _analysis(summary="р & <т> " * 800, risks=["риск " * 300] * 3))
    assert len(msg) <= TELEGRAM_MAX_MESSAGE
    assert 'href="https://x/1"' in msg
    assert msg.count("<b>") == msg.count("</b>")


def test_excel_contains_description_column(tmp_path):
    db = TenderDatabase(tmp_path / "t.sqlite3")
    t = Tender(platform="fabrikant", external_id="9", title="Насос", url="https://x/9",
               description="Поставка насосов")
    db.save_tender(t)
    tid = db.get_tender_id(t.unique_key)
    path = export_tenders_to_excel(db, tmp_path / "o.xlsx", tender_ids=[tid])
    ws = load_workbook(path)["Тендеры"]
    headers = [c.value for c in ws[1]]
    assert headers[-1] == "Описание"
    assert ws.cell(row=2, column=headers.index("Описание") + 1).value == "Поставка насосов"
    assert ws.cell(row=2, column=1).value == "Фабрикант"
