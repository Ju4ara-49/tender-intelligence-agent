"""Regression-тесты отображения дат в Excel-отчёте.

База хранит даты в UTC (контракт Tender.to_utc), но пользовательский отчёт
обязан показывать московское время: дедлайны закупок РФ публикуются в МСК.
"""
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from src.export.excel import MOSCOW_TZ, _excel_datetime, export_tenders_to_excel
from src.models.tender import Tender
from src.storage.database import TenderDatabase


def test_excel_datetime_converts_utc_to_moscow_wall_time():
    utc_value = datetime(2026, 9, 30, 21, 0, tzinfo=timezone.utc)
    # 21:00 UTC == 00:00 МСК следующего дня.
    assert _excel_datetime(utc_value) == datetime(2026, 10, 1, 0, 0)


def test_excel_datetime_handles_none():
    assert _excel_datetime(None) is None


def test_excel_deadline_cell_shows_moscow_time(tmp_path):
    db = TenderDatabase(tmp_path / "msk.db")
    deadline = datetime(2026, 10, 1, 12, 0, tzinfo=timezone.utc)  # 15:00 МСК
    tender = Tender(
        platform="test",
        external_id="msk-display-1",
        title="Проверка МСК-времени",
        url="https://example.test/msk-1",
        deadline=deadline,
    )
    tender_id = db.save_tender(tender)

    output = export_tenders_to_excel(db, tmp_path / "msk.xlsx", tender_ids=[tender_id])
    ws = load_workbook(output)["Тендеры"]

    assert ws["I2"].value == datetime(2026, 10, 1, 15, 0)
    assert ws["I2"].number_format == "dd.mm.yyyy"


def test_days_left_uses_moscow_dates_not_utc(tmp_path):
    """Regression: дедлайн 00:30 МСК — это 21:30 UTC предыдущего дня.
    Старый расчёт по UTC-датам занижал «Осталось дней» на единицу."""
    now_msk = datetime.now(MOSCOW_TZ)
    deadline_msk = (now_msk + timedelta(days=3)).replace(hour=0, minute=30, second=0, microsecond=0)
    assert deadline_msk.date() == (now_msk + timedelta(days=3)).date()

    db = TenderDatabase(tmp_path / "msk-days.db")
    tender = Tender(
        platform="test",
        external_id="msk-days-1",
        title="Осталось дней по МСК",
        url="https://example.test/msk-days-1",
        deadline=deadline_msk,
    )
    tender_id = db.save_tender(tender)

    output = export_tenders_to_excel(db, tmp_path / "msk-days.xlsx", tender_ids=[tender_id])
    ws = load_workbook(output)["Тендеры"]

    assert ws["I2"].value == datetime.combine(deadline_msk.date(), datetime.min.time()).replace(hour=0, minute=30)
    assert ws["J2"].value == 3