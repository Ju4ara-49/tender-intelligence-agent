from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

from src.models.tender import Tender
from src.orchestrator import Orchestrator
from src.telegram_settings import TenderCriteria


class _DetailCollector:
    platform = "rts_tender"

    def get_details(self, external_id: str) -> Tender:
        return Tender(
            platform=self.platform,
            external_id=external_id,
            title="Поставка запасных частей — детали",
            url="https://example.test/tender/1",
            description="Расширенное описание",
            price=1_500_000,
            start_date=datetime(2026, 9, 12, 9, 0),
            end_date=datetime(2026, 9, 20, 18, 0),
            published_at=datetime(2026, 9, 11, 8, 0),
            deadline=datetime(2026, 9, 20, 18, 0),
            region="Санкт-Петербург",
            customer="ООО Ромашка",
            law_type="223-ФЗ",
            advance_required=True,
            advance_percent=30,
            postpayment_days=15,
            application_security_percent=2,
            contract_security_percent=5,
            raw_data={"items": ["подшипник"]},
        )


def _tender(**overrides) -> Tender:
    values = {
        "platform": "rts_tender",
        "external_id": "1",
        "title": "Тестовый тендер",
        "url": "https://example.test/tender/1",
        "price": 1_000_000,
        "deadline": datetime.now(timezone.utc) + timedelta(days=10),
        "advance_required": True,
        "advance_percent": 30,
        "postpayment_days": 15,
        "application_security_percent": 2,
        "contract_security_percent": 5,
    }
    values.update(overrides)
    return Tender(**values)


def test_enrichment_copies_all_common_detail_fields() -> None:
    base = _tender(start_date=None, end_date=None, deadline=None, published_at=None, advance_required=False, advance_percent=None, postpayment_days=None, application_security_percent=None, contract_security_percent=None)
    enriched, loaded = Orchestrator.__new__(Orchestrator)._enrich_tender(_DetailCollector(), base)

    assert loaded is True
    assert enriched.start_date is not None
    assert enriched.end_date is not None
    assert enriched.deadline == enriched.end_date
    assert enriched.advance_required is True
    assert enriched.advance_percent == 30
    assert enriched.postpayment_days == 15
    assert enriched.application_security_percent == 2
    assert enriched.contract_security_percent == 5
    assert enriched.customer == "ООО Ромашка"
    assert enriched.region == "Санкт-Петербург"


def test_detail_values_replace_stale_search_values() -> None:
    base = _tender(
        start_date=datetime(2026, 9, 10, 9, 0),
        end_date=datetime(2026, 9, 18, 18, 0),
        deadline=datetime(2026, 9, 18, 18, 0),
        region="Москва",
        customer="Старый заказчик",
    )
    enriched, _ = Orchestrator.__new__(Orchestrator)._enrich_tender(_DetailCollector(), base)

    assert enriched.end_date == datetime(2026, 9, 20, 15, 0, tzinfo=timezone.utc)
    assert enriched.deadline == enriched.end_date
    assert enriched.region == "Санкт-Петербург"
    assert enriched.customer == "ООО Ромашка"


def test_max_postpayment_rejects_missing_value() -> None:
    tender = _tender(postpayment_days=None)
    criteria = TenderCriteria(max_postpayment_days=30)
    assert Orchestrator._passes_criteria(tender, criteria) == (False, "max_postpayment_days")


def test_max_security_rejects_missing_value() -> None:
    tender = _tender(application_security_percent=None, contract_security_percent=None)
    criteria = TenderCriteria(max_application_security_percent=5, max_contract_security_percent=10)
    assert Orchestrator._passes_criteria(tender, criteria) == (False, "max_application_security_percent")


def test_max_postpayment_and_security_accept_values_inside_limits() -> None:
    tender = _tender(postpayment_days=30, application_security_percent=5, contract_security_percent=10)
    criteria = TenderCriteria(max_postpayment_days=30, max_application_security_percent=5, max_contract_security_percent=10)
    assert Orchestrator._passes_criteria(tender, criteria) == (True, "")


def test_platform_worker_count_honors_concurrency_alias_and_bounds():
    assert Orchestrator._platform_worker_count({"concurrency": 2}, 6) == 2
    assert Orchestrator._platform_worker_count({"platform_workers": 4, "concurrency": 2}, 6) == 4
    assert Orchestrator._platform_worker_count({"concurrency": 99}, 6) == 6
    assert Orchestrator._platform_worker_count({"concurrency": 0}, 6) == 1
    assert Orchestrator._platform_worker_count({"concurrency": "bad"}, 6) == 6


def test_keyword_filter_search_documents_controls_text_scope():
    """When search_documents=False, document text is excluded from matching."""
    from src.filters.keyword_filter import KeywordFilter

    tender = _tender(
        title="Закупка оборудования",
        description="Техническое задание",
        raw_data={"document_search_text": "подшипники SKF подходят"},
    )
    full_filt = KeywordFilter(["подшипники"], [], min_text_length=1, search_documents=True)
    short_filt = KeywordFilter(["подшипники"], [], min_text_length=1, search_documents=False)
    assert full_filt.matches_strict(tender)
    assert not short_filt.matches_strict(tender)


def test_matches_law_type_shorthand():
    criteria = TenderCriteria(law_type="44")
    assert Orchestrator._passes_criteria(_tender(law_type="44-ФЗ"), criteria) == (True, "")
    assert Orchestrator._passes_criteria(_tender(law_type="223-ФЗ"), criteria)[0] is False


def test_matches_customer_contains():
    criteria = TenderCriteria(customer="Ромашка")
    assert Orchestrator._passes_criteria(_tender(customer="ООО Ромашка"), criteria) == (True, "")
    assert Orchestrator._passes_criteria(_tender(customer="ООО Солнечко"), criteria)[0] is False


def test_matches_customer_inn_exact():
    criteria = TenderCriteria(customer_inn="7701234567")
    assert Orchestrator._passes_criteria(_tender(customer_inn="7701234567"), criteria) == (True, "")
    assert Orchestrator._passes_criteria(_tender(customer_inn="7709999999"), criteria)[0] is False


def test_excel_export_uses_post_filter_result_ids_not_diagnostic_persistence_ids():
    source = Path(Orchestrator.__module__.replace(".", "/") + ".py").read_text(encoding="utf-8")
    assert "export_tender_ids: list[int] = []" in source
    assert "export_tender_ids.append(tender_id)" in source
    assert "tender_ids=export_tender_ids" in source
    assert "tender_ids=current_run_tender_ids" not in source

def test_notification_dedup_does_not_skip_ai_or_lifecycle_progression():
    source = Path(Orchestrator.__module__.replace(".", "/") + ".py").read_text(encoding="utf-8")
    analyze_pos = source.index("analysis = self.analyzer.analyze(")
    dedup_pos = source.index(
        'if self.notification_state.was_notified(tender, recipient_key=recipient_key):',
        analyze_pos,
    )
    lifecycle_pos = source.index(
        "self._advance_lifecycle(tender, TenderLifecycleStatus.RELEVANT)",
        analyze_pos,
    )
    assert analyze_pos < lifecycle_pos < dedup_pos


def test_search_profile_deterministic_end_to_end_to_excel(monkeypatch, tmp_path):
    from openpyxl import load_workbook

    from src.models.tender import TenderAnalysis
    from src.profiles import SearchProfileStore
    from src.settings import AppSettings
    from src.tenderplan import TenderLifecycleStatus, application_task_id

    class FakeCollector:
        platform = "eis"
        deadline = datetime(2030, 1, 20, 12, 0, tzinfo=timezone.utc)

        def __init__(self):
            self.price = 100000.0

        def search(self, keywords, since=None):
            return [Tender(
                platform=self.platform,
                external_id="E2E-1",
                title="Поставка подшипников",
                url="https://example.test/e2e-1",
                description="Подшипники для оборудования",
            )]

        def get_details(self, external_id):
            return Tender(
                platform=self.platform,
                external_id=external_id,
                title="Поставка подшипников",
                url="https://example.test/e2e-1",
                description="Подшипники для оборудования",
                price=self.price,
                deadline=self.deadline,
                region="Москва",
                customer="ООО Е2Е",
                customer_inn="7701234567",
                law_type="44-ФЗ",
                advance_required=True,
                advance_percent=20,
                postpayment_days=15,
                application_security_percent=2,
                contract_security_percent=10,
                okpd2_codes=["26.30.11"],
                procurement_type="commercial",
            )

    class FakeAnalyzer:
        def analyze(self, tender, *, search_documents=True):
            return TenderAnalysis(85, "Подходит", "participate")

    class FakeNotifier:
        def __init__(self):
            self.sent = []

        def send_tender_alert(self, tender, analysis, chat_id=None):
            self.sent.append((tender.unique_key, chat_id))
            return True

    db_path = tmp_path / "e2e.db"
    output_dir = tmp_path / "output"
    settings = AppSettings(
        config={
            "storage": {"database_path": str(db_path)},
            "export": {"output_dir": str(output_dir)},
            "filters": {"min_text_length": 1},
            "search": {"platform_workers": 1},
            "notifications": {"telegram": {"dry_run_when_no_token": True}},
        },
        keywords={"include": ["подшипники"]},
    )
    collector = FakeCollector()
    monkeypatch.setattr("src.orchestrator.get_enabled_collectors", lambda config, enabled_platforms=None: [collector])

    orchestrator = Orchestrator(settings)
    orchestrator.analyzer = FakeAnalyzer()
    notifier = FakeNotifier()
    orchestrator.notifier = notifier
    SearchProfileStore(orchestrator.db).create(
        "user-e2e",
        name="Е2Е профиль",
        keywords=["подшипники"],
        platforms=["eis"],
        regions=["Москва"],
        min_price=50000,
        max_price=200000,
        advance_required=True,
        min_advance_percent=10,
        max_postpayment_days=30,
        min_submission_days=7,
        max_application_security_percent=5,
        max_contract_security_percent=20,
        okpd2_codes=["26.30"],
        procurement_types=["commercial"],
    )

    first = orchestrator.run_cycle_for_user("user-e2e")
    assert first[0]["new"] == 1
    assert first[0]["analyzed"] == 1
    assert first[0]["notified"] == 1
    assert orchestrator.lifecycle_store.get("eis:E2E-1") is TenderLifecycleStatus.SHORTLISTED
    assert orchestrator.task_store.get(application_task_id("eis:E2E-1", "user-e2e"), user_id="user-e2e") is not None
    first_excel = sorted(output_dir.glob("search_*.xlsx"))[-1]
    first_sheet = load_workbook(first_excel)["Тендеры"]
    assert first_sheet.max_row == 2
    assert first_sheet["C2"].value == "Поставка подшипников"

    second = orchestrator.run_cycle_for_user("user-e2e")
    assert second[0]["new"] == 0
    assert second[0]["skipped_duplicate"] == 1
    assert len(notifier.sent) == 1

    collector.price = 120000.0
    third = orchestrator.run_cycle_for_user("user-e2e")
    assert third[0]["new"] == 0
    assert third[0]["notified"] == 1
    assert len(notifier.sent) == 2
    assert orchestrator.db.get_tender("eis:E2E-1").price == 120000.0
    tender_id = orchestrator.db.get_tender_id("eis:E2E-1")
    assert len(orchestrator.db.get_tender_history(tender_id)) == 2


class _FakeDocument:
    """Minimal stand-in for a TenderDocument ingest result."""

    def __init__(self, url: str, text: str, status: str):
        import hashlib

        self.url = url
        self.extracted_text = text
        self.extraction_status = status
        self.version = 1
        self.sha256 = hashlib.sha256(url.encode()).hexdigest()
        self.diagnostics = ""


class _FakeIngestResult:
    def __init__(self, document: _FakeDocument, downloaded: bool = True):
        self.document = document
        self.downloaded = downloaded


class _FakeDocumentIngestor:
    """Deterministic document ingestor: URL ending in '#fail' -> failed extraction."""

    def __init__(self):
        self.calls: list[str] = []

    def ingest(self, *, tender_key: str, url: str, filename: str = "", content_type: str = ""):
        self.calls.append(url)
        if url.endswith("#fail"):
            return _FakeIngestResult(_FakeDocument(url, "", "failed"))
        text = f"Техническое задание: поставка подшипников SKF для {filename or url}"
        return _FakeIngestResult(_FakeDocument(url, text, "extracted"))


class _DocumentCollector:
    """Discovery title has NO keyword; the keyword lives only in documents."""

    platform = "eis"

    def __init__(self, documents: list[dict[str, str]]):
        self.documents = documents

    def search(self, keywords, since=None):
        return [Tender(
            platform=self.platform,
            external_id="DOC-1",
            title="Закупка оборудования для цеха",
            url="https://example.test/doc-1",
        )]

    def get_details(self, external_id: str) -> Tender:
        return Tender(
            platform=self.platform,
            external_id=external_id,
            title="Закупка оборудования для цеха",
            url="https://example.test/doc-1",
            description="Подробное описание потребности",
            price=300000.0,
            deadline=datetime(2030, 1, 20, 12, 0, tzinfo=timezone.utc),
            region="Москва",
            customer="ООО Документ",
            documents=list(self.documents),
        )


def _run_document_cycle(monkeypatch, tmp_path, documents, search_documents=True):
    from src.models.tender import TenderAnalysis
    from src.profiles import SearchProfileStore
    from src.settings import AppSettings

    class FakeAnalyzer:
        def analyze(self, tender, *, search_documents=True):
            return TenderAnalysis(90, "Подшипники в документе", "participate")

    class FakeNotifier:
        def __init__(self):
            self.sent = []

        def send_tender_alert(self, tender, analysis, chat_id=None):
            self.sent.append(tender.unique_key)
            return True

    settings = AppSettings(
        config={
            "storage": {"database_path": str(tmp_path / "doc.db")},
            "export": {"output_dir": str(tmp_path / "out")},
            "filters": {"min_text_length": 1},
            "search": {"platform_workers": 1},
            "notifications": {"telegram": {"dry_run_when_no_token": True}},
        },
        keywords={"include": ["подшипник"]},
    )
    collector = _DocumentCollector(documents)
    ingestor = _FakeDocumentIngestor()
    monkeypatch.setattr(
        "src.orchestrator.get_enabled_collectors",
        lambda config, enabled_platforms=None: [collector],
    )
    orchestrator = Orchestrator(settings)
    orchestrator.document_ingestor = ingestor
    orchestrator.analyzer = FakeAnalyzer()
    notifier = FakeNotifier()
    orchestrator.notifier = notifier
    SearchProfileStore(orchestrator.db).create(
        "user-doc",
        name="Документный профиль",
        keywords=["подшипник"],
        platforms=["eis"],
        document_search=search_documents,
        max_application_security_percent=None,
        max_contract_security_percent=None,
    )
    result = orchestrator.run_cycle_for_user("user-doc")
    return orchestrator, result[0], notifier, ingestor


def test_document_search_finds_keyword_only_in_document(monkeypatch, tmp_path):
    """Keyword exists only in an extracted document -> tender is found and indexed."""
    documents = [{"name": "tz.pdf", "url": "https://example.test/tz.pdf"}]
    orchestrator, stats, notifier, ingestor = _run_document_cycle(monkeypatch, tmp_path, documents)

    assert stats["found"] == 1
    assert stats["new"] == 1
    assert stats["notified"] == 1
    assert stats["analyzed"] == 1
    assert ingestor.calls == ["https://example.test/tz.pdf"]
    stored = orchestrator.db.get_tender("eis:DOC-1")
    assert "подшипник" in stored.full_text

    second = orchestrator.run_cycle_for_user("user-doc")
    assert second[0]["skipped_duplicate"] == 1
    assert second[0]["notified"] == 0
    assert len(notifier.sent) == 1


def test_document_search_off_does_not_match_document_keyword(monkeypatch, tmp_path):
    """With document_search=False the same tender must NOT match."""
    documents = [{"name": "tz.pdf", "url": "https://example.test/tz.pdf"}]
    _, stats, _, _ = _run_document_cycle(monkeypatch, tmp_path, documents, search_documents=False)
    assert stats["filtered"] == 0
    assert stats["keyword_excluded"] == 1


def test_failed_document_is_not_claimed_as_indexed(monkeypatch, tmp_path):
    """A failed extraction provides no text; only the good document is searchable."""
    documents = [
        {"name": "tz.pdf", "url": "https://example.test/tz.pdf#fail"},
        {"name": "spec.pdf", "url": "https://example.test/spec.pdf"},
    ]
    orchestrator, stats, _, _ = _run_document_cycle(monkeypatch, tmp_path, documents)

    assert stats["documents_discovered"] == 2
    stored = orchestrator.db.get_tender("eis:DOC-1")
    assert "spec.pdf" in stored.raw_data.get("document_contents", "")
    assert "#fail" not in stored.raw_data.get("document_contents", "")
    versions = stored.raw_data.get("document_versions", [])
    failed = [v for v in versions if v["extraction_status"] == "failed"]
    assert len(failed) == 1
    assert failed[0]["url"] == "https://example.test/tz.pdf#fail"


def test_tender_without_documents_still_passes_pipeline(monkeypatch, tmp_path):
    """No documents: keyword never matches -> no save-notification, no ingest calls."""
    orchestrator, stats, notifier, ingestor = _run_document_cycle(monkeypatch, tmp_path, [])
    assert stats["documents_discovered"] == 0
    assert ingestor.calls == []
    assert notifier.sent == []
    # Keyword "подшипник" is absent everywhere -> strict filter excludes the tender.
    assert stats["filtered"] == 0
    assert stats["keyword_excluded"] == 1
    assert stats["notified"] == 0
