"""Безопасное извлечение текста из документов тендера без облачного AI."""
from __future__ import annotations

import html
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from openpyxl import load_workbook

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - dependency is declared in requirements
    PdfReader = None


MAX_ARCHIVE_FILES = 500
MAX_ARCHIVE_UNCOMPRESSED = 200 * 1024 * 1024
MAX_FILE_BYTES = 50 * 1024 * 1024
SUPPORTED_TEXT_EXTENSIONS = {".txt", ".csv", ".xml", ".html", ".htm", ".json"}
SUPPORTED_OFFICE_EXTENSIONS = {".docx", ".xlsx", ".xlsm"}
SUPPORTED_ARCHIVES = {".zip"}
SUPPORTED_PDF = {".pdf"}


@dataclass(frozen=True)
class DocumentText:
    path: str
    text: str
    mime_hint: str = ""


@dataclass(frozen=True)
class DocumentHit:
    path: str
    keyword: str
    snippet: str
    start: int
    end: int


class UnsafeArchiveError(ValueError):
    """Архив содержит опасный путь или превышает лимиты."""


class DocumentIntelligence:
    """Извлечение, нормализация и локальный поиск по документам и архивам."""

    def __init__(
        self,
        max_file_bytes: int = MAX_FILE_BYTES,
        max_archive_files: int = MAX_ARCHIVE_FILES,
        max_archive_uncompressed: int = MAX_ARCHIVE_UNCOMPRESSED,
    ) -> None:
        self.max_file_bytes = max_file_bytes
        self.max_archive_files = max_archive_files
        self.max_archive_uncompressed = max_archive_uncompressed

    @staticmethod
    def normalize(text: str) -> str:
        text = html.unescape(text or "")
        text = text.replace("\x00", " ")
        text = re.sub(r"[\t\r ]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def extract_file(self, path: Path, virtual_path: str | None = None) -> list[DocumentText]:
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        if path.stat().st_size > self.max_file_bytes:
            raise ValueError(f"Файл слишком большой: {path}")
        name = virtual_path or path.name
        ext = path.suffix.lower()
        if ext in SUPPORTED_ARCHIVES:
            return self._extract_archive(path.read_bytes(), name)
        return self._extract_bytes(path.read_bytes(), name, ext)

    def extract_bytes(self, data: bytes, filename: str) -> list[DocumentText]:
        if len(data) > self.max_file_bytes:
            raise ValueError(f"Файл слишком большой: {filename}")
        return self._extract_bytes(data, filename, Path(filename).suffix.lower())

    def _extract_bytes(self, data: bytes, name: str, ext: str) -> list[DocumentText]:
        if ext in SUPPORTED_TEXT_EXTENSIONS:
            return [DocumentText(name, self.normalize(data.decode("utf-8", errors="replace")), ext)]
        if ext in SUPPORTED_PDF:
            return [DocumentText(name, self._pdf(data), "application/pdf")]
        if ext in SUPPORTED_OFFICE_EXTENSIONS:
            if ext in {".xlsx", ".xlsm"}:
                return [DocumentText(name, self._xlsx(data), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")]
            return [DocumentText(name, self._docx(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")]
        if ext in SUPPORTED_ARCHIVES:
            return self._extract_archive(data, name)
        return []

    def _pdf(self, data: bytes) -> str:
        if PdfReader is None:
            raise RuntimeError("Для PDF нужен пакет pypdf")
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return self.normalize("\n\n".join(pages))

    def _docx(self, data: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            try:
                xml = archive.read("word/document.xml")
            except KeyError:
                return ""
        root = ElementTree.fromstring(xml)
        texts = []
        for node in root.iter():
            if node.tag.endswith("}t") and node.text:
                texts.append(node.text)
            elif node.tag.endswith("}tab"):
                texts.append("\t")
            elif node.tag.endswith("}br"):
                texts.append("\n")
        return self.normalize("".join(texts))

    def _xlsx(self, data: bytes) -> str:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        try:
            for sheet in workbook.worksheets:
                lines.append(f"[Лист: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        lines.append(" | ".join(values))
        finally:
            workbook.close()
        return self.normalize("\n".join(lines))

    def _extract_archive(self, data: bytes, name: str) -> list[DocumentText]:
        results: list[DocumentText] = []
        total = 0
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = [info for info in archive.infolist() if not info.is_dir()]
            if len(infos) > self.max_archive_files:
                raise UnsafeArchiveError("Архив содержит слишком много файлов")
            for info in infos:
                self._validate_archive_member(info.filename)
                total += info.file_size
                if total > self.max_archive_uncompressed:
                    raise UnsafeArchiveError("Распакованный размер архива превышает лимит")
                member = archive.read(info)
                if len(member) > self.max_file_bytes:
                    continue
                virtual = f"{name}!/{info.filename}"
                try:
                    results.extend(self._extract_bytes(member, virtual, Path(info.filename).suffix.lower()))
                except (ValueError, zipfile.BadZipFile, UnicodeError):
                    continue
        return results

    @staticmethod
    def _validate_archive_member(name: str) -> None:
        normalized = name.replace("\\", "/")
        path = PurePosixPath(normalized)
        if path.is_absolute() or ".." in path.parts:
            raise UnsafeArchiveError(f"Опасный путь в архиве: {name}")
        if re.match(r"^[A-Za-z]:", normalized):
            raise UnsafeArchiveError(f"Абсолютный путь в архиве: {name}")

    def search(self, documents: list[DocumentText], keywords: list[str], context: int = 100) -> list[DocumentHit]:
        hits: list[DocumentHit] = []
        for document in documents:
            text_lower = document.text.casefold()
            for keyword in keywords:
                needle = str(keyword).strip()
                if not needle:
                    continue
                start = 0
                needle_lower = needle.casefold()
                while True:
                    position = text_lower.find(needle_lower, start)
                    if position < 0:
                        break
                    left = max(0, position - context)
                    right = min(len(document.text), position + len(needle) + context)
                    snippet = document.text[left:right].replace("\n", " ").strip()
                    hits.append(DocumentHit(document.path, needle, snippet, position, position + len(needle)))
                    start = position + max(1, len(needle))
                    if len(hits) >= 500:
                        return hits
        return hits

    def analyze_paths(self, paths: list[Path], keywords: list[str]) -> tuple[list[DocumentText], list[DocumentHit]]:
        documents: list[DocumentText] = []
        for path in paths:
            documents.extend(self.extract_file(Path(path)))
        return documents, self.search(documents, keywords)
